import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Colors (Crisp Studio Light Theme - Professional Wide Audience Style)
BG_CANVAS = RGBColor(248, 250, 252)       # #F8FAFC (Soft Light Studio Canvas)
SURFACE_CARD = RGBColor(255, 255, 255)    # #FFFFFF (Pure Crisp White)
BORDER_CARD = RGBColor(226, 232, 240)     # #E2E8F0 (Delicate Slate Border)
TEXT_MAIN = RGBColor(15, 23, 42)          # #0F172A (Deep Obsidian Ink)
TEXT_MUTED = RGBColor(71, 85, 105)        # #475569 (Slate Muted Gray)
ACCENT_INDIGO = RGBColor(67, 56, 202)     # #4338CA (Royal Indigo Accent)
ACCENT_PURPLE = RGBColor(124, 58, 237)    # #7C3AED (Qwen Violet Highlight)
ACCENT_EMERALD = RGBColor(4, 120, 87)     # #047857 (Forest Emerald Stat)
ACCENT_AMBER = RGBColor(192, 86, 33)      # #C05621 (Warm Amber Alert)
ACCENT_ROSE = RGBColor(225, 29, 72)       # #E11D48 (Rose Red Alert)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_CANVAS

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.35))
        tf = tag_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.name = FONT_BODY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_HEADING
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN

        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.4))
            tf = sub_box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = subtitle_text
            p.font.name = FONT_BODY
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, bg_color=SURFACE_CARD, border_color=BORDER_CARD):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide (Broad Audience Version)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    add_card(s1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), bg_color=SURFACE_CARD, border_color=BORDER_CARD)

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(1.35), Inches(3.4), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(238, 242, 255)
    badge.line.color.rgb = ACCENT_INDIGO
    badge.line.width = Pt(1)
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.text = "STRATEGIC RESEARCH BRIEFING"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    title_box = s1.shapes.add_textbox(Inches(1.3), Inches(2.0), Inches(10.733), Inches(1.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Augmenting Research & Analytics with AI"
    p.font.name = FONT_HEADING
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN

    p2 = tf.add_paragraph()
    p2.text = "Leveraging Qwen 2.5, DeepSeek-R1, TabFM & Streamlit for Data-Driven Insights"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(17)
    p2.font.color.rgb = ACCENT_INDIGO

    stats_box = s1.shapes.add_textbox(Inches(1.3), Inches(4.2), Inches(10.733), Inches(2.1))
    tf = stats_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Core Strategic Deliverables:"
    p.font.name = FONT_BODY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD

    bullets = [
        "Qwen Multilingual AI: Zero-shot coding across Hindi & English field notes",
        "70% faster turnaround on qualitative synthesis and evidence extraction",
        "TabFM Econometric Forecasting: Merging observational field narratives with statistics",
        "Interactive Streamlit Dashboard: Real-time analytical portal with full evidence auditability"
    ]
    for bullet in bullets:
        pb = tf.add_paragraph()
        pb.text = f"•  {bullet}"
        pb.font.name = FONT_BODY
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 2: Executive Summary (3 Pillars)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Executive Summary", "Why AI in Our Research & Analytics Operations?", "Transforming qualitative field data into actionable policy decisions with speed and econometric rigor.")

    pillars = [
        ("70%", "Speed & Efficiency", "Reduces manual reading and line-by-line coding. Qwen 2.5 & DeepSeek synthesize thousands of Hindi/English field notes in minutes."),
        ("Qual + Quant", "Dual-Engine Intelligence", "Combines qualitative theme discovery (Qwen) with econometric predictive modeling (TabFM) to forecast adoption and participation rates."),
        ("$0 Extra", "Zero Added Infra Cost", "Operates on open-weight architectures, direct terminal shell automation, and lightweight Streamlit interfaces with zero server lock-in.")
    ]

    card_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    card_top = Inches(2.0)
    card_h = Inches(4.8)

    for i, (stat, title, desc) in enumerate(pillars):
        x = start_x + i * (card_w + gap)
        add_card(s2, x, card_top, card_w, card_h)

        tb = s2.shapes.add_textbox(x + Inches(0.3), card_top + Inches(0.4), card_w - Inches(0.6), card_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stat
        p.font.name = FONT_BODY
        p.font.size = Pt(38)
        p.font.bold = True
        p.font.color.rgb = ACCENT_PURPLE if i == 0 else (ACCENT_INDIGO if i == 1 else ACCENT_EMERALD)

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: Current Workflow Challenges vs AI Solution
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Workflow Transformation", "Bridging Current Bottlenecks with AI Capability", "Shifting our research team from manual sorting to strategic analysis.")

    col_w = Inches(5.66)
    c1_x = Inches(0.8)
    c2_x = Inches(6.86)
    c_top = Inches(2.0)
    c_h = Inches(4.8)

    add_card(s3, c1_x, c_top, col_w, c_h, border_color=ACCENT_ROSE)
    tb1 = s3.shapes.add_textbox(c1_x + Inches(0.3), c_top + Inches(0.3), col_w - Inches(0.6), c_h - Inches(0.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "TRADITIONAL RESEARCH BOTTLENECKS"
    p.font.name = FONT_BODY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE

    t_points = [
        ("Manual Text Coding", "Hundreds of hours spent reading and categorizing Hindi/English field notes line-by-line."),
        ("Siloed Data", "Qualitative observer notes and quantitative survey numbers analyzed in separate streams."),
        ("Slow Turnaround", "Weeks between field observation visits and final research briefings."),
        ("No Predictive Insight", "Historical reporting only; unable to predict participant drop-offs before intervention.")
    ]

    for title, detail in t_points:
        p_sub = tf1.add_paragraph()
        p_sub.text = f"\n❌  {title}"
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(14)
        p_sub.font.bold = True
        p_sub.font.color.rgb = TEXT_MAIN

        p_det = tf1.add_paragraph()
        p_det.text = f"     {detail}"
        p_det.font.name = FONT_BODY
        p_det.font.size = Pt(12)
        p_det.font.color.rgb = TEXT_MUTED

    add_card(s3, c2_x, c_top, col_w, c_h, border_color=ACCENT_EMERALD)
    tb2 = s3.shapes.add_textbox(c2_x + Inches(0.3), c_top + Inches(0.3), col_w - Inches(0.6), c_h - Inches(0.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "AI-AUGMENTED RESEARCH PIPELINE"
    p.font.name = FONT_BODY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD

    ai_points = [
        ("Qwen Multilingual AI", "Extracts themes & zero-shot codes Hindi & English field notes automatically."),
        ("Closed-Loop Fusion", "Qualitative themes convert to feature vectors directly for TabFM predictive forecasting."),
        ("70% Faster Delivery", "Draft summaries delivered in hours; researchers focus on validation & nuance."),
        ("Proactive Forecasting", "Early-warning predictions highlight at-risk schools before full program rollout.")
    ]

    for title, detail in ai_points:
        p_sub = tf2.add_paragraph()
        p_sub.text = f"\n✅  {title}"
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(14)
        p_sub.font.bold = True
        p_sub.font.color.rgb = TEXT_MAIN

        p_det = tf2.add_paragraph()
        p_det.text = f"     {detail}"
        p_det.font.name = FONT_BODY
        p_det.font.size = Pt(12)
        p_det.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: The Integrated AI Tool Stack
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "AI Tool Ecosystem", "The Core Components of Our AI Architecture", "Specialized open-weight tools working seamlessly together to power analysis.")

    tools = [
        ("Qwen 2.5", "Multilingual Text Copilot", "Organizes multilingual field notes, extracts themes from Hindi & English observer notes, groups passages, and drafts summary reports.", ACCENT_PURPLE),
        ("DeepSeek-R1", "Reasoning & Audit Engine", "Performs deep cross-checks, compares findings across projects, flags contradictions, and generates interpretive hypotheses.", ACCENT_INDIGO),
        ("TabFM", "Predictive Foundation Model", "Uses spreadsheet data and qualitative feature vectors to forecast outcomes (e.g. participation drop-off) with statistical rigor.", ACCENT_AMBER),
        ("Streamlit & Antigravity", "EDS & CRO Portals", "Provides interactive analytics dashboards (EDS: cm-rise-eds-dashboard.streamlit.app | CRO: peepul-mp-cpd-cro-dashboard.streamlit.app) and automates terminal execution pipelines.", ACCENT_EMERALD)
    ]

    card_w = Inches(5.66)
    card_h = Inches(2.25)
    coords = [
        (Inches(0.8), Inches(2.0)),
        (Inches(6.86), Inches(2.0)),
        (Inches(0.8), Inches(4.55)),
        (Inches(6.86), Inches(4.55))
    ]

    for i, (name, role, desc, color) in enumerate(tools):
        x, y = coords[i]
        add_card(s4, x, y, card_w, card_h, border_color=BORDER_CARD)

        tb = s4.shapes.add_textbox(x + Inches(0.3), y + Inches(0.25), card_w - Inches(0.6), card_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = name
        p.font.name = FONT_HEADING
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = color

        p_r = tf.add_paragraph()
        p_r.text = role.upper()
        p_r.font.name = FONT_BODY
        p_r.font.size = Pt(10)
        p_r.font.bold = True
        p_r.font.color.rgb = TEXT_MUTED

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 5: Closed-Loop AI Pipeline (6 Steps)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Integrated Workflow", "Closed-Loop Pipeline: From Field Notes to Analytical Briefing", "A repeatable, automated workflow bridging text coding with quantitative forecasting.")

    steps = [
        ("1. Multilingual Notes", "Field notes in Hindi & English collected from school visits."),
        ("2. Qwen Coding", "Qwen codes text; DeepSeek audits themes & contradictions."),
        ("3. Feature Vector", "Extracted themes converted into quantitative feature vectors."),
        ("4. TabFM Predict", "Predictive modeling estimates risk & adoption rates."),
        ("5. Data & Charts", "Terminal & AGY CLI automate chart creation & audit trail."),
        ("6. Analytics Portal", "Results deployed to Streamlit dashboard for team decision.")
    ]

    step_w = Inches(1.8)
    step_gap = Inches(0.18)
    start_x = Inches(0.8)
    y_pos = Inches(2.2)
    h_pos = Inches(4.5)

    for i, (title, desc) in enumerate(steps):
        x = start_x + i * (step_w + step_gap)
        color = ACCENT_PURPLE if i == 1 else (ACCENT_INDIGO if i % 2 == 0 else ACCENT_AMBER)
        add_card(s5, x, y_pos, step_w, h_pos, border_color=BORDER_CARD)

        tb = s5.shapes.add_textbox(x + Inches(0.15), y_pos + Inches(0.2), step_w - Inches(0.3), h_pos - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MAIN

    adv_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.4))
    adv_card.fill.solid()
    adv_card.fill.fore_color.rgb = RGBColor(238, 242, 255)
    adv_card.line.color.rgb = ACCENT_INDIGO
    adv_card.line.width = Pt(1)
    tf = adv_card.text_frame
    p = tf.paragraphs[0]
    p.text = "KEY ADVANTAGE: 70% faster turnaround, standardized coding consistency, and evidence-grounded forecasting."
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_BODY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    # =========================================================================
    # SLIDE 6: Empirical Proof 1: EDS Teacher Readiness
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Empirical Proof #1", "EDS Study: Identifying Intervention Gaps Instantly", "Demonstrating Qwen AI analysis on teacher training and material readiness (N = 60).")

    add_card(s6, Inches(0.8), Inches(2.0), Inches(4.5), Inches(4.8), border_color=ACCENT_AMBER)
    tb = s6.shapes.add_textbox(Inches(1.1), Inches(2.4), Inches(3.9), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "19 of 60"
    p.font.name = FONT_BODY
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    p_lbl = tf.add_paragraph()
    p_lbl.text = "Teachers Received NO Training Materials\n"
    p_lbl.font.name = FONT_HEADING
    p_lbl.font.size = Pt(16)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = TEXT_MAIN

    p_d = tf.add_paragraph()
    p_d.text = "Study Sample: N = 60 teachers across participating EDS districts.\n\nQwen parsed multilingual qualitative feedback from field observers, instantly categorizing that nearly one-third lacked physical guides.\n\n📄 Referenced Official Report: EDS_Report_TabFM.pdf (Qualitative Field Insights: Observation Sheet Analysis)"
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(11)
    p_d.font.color.rgb = TEXT_MUTED

    add_card(s6, Inches(5.6), Inches(2.0), Inches(6.933), Inches(4.8))
    tb2 = s6.shapes.add_textbox(Inches(5.9), Inches(2.3), Inches(6.333), Inches(4.2))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "INSIGHTS & ACTIONABLE RESPONSE"
    p.font.name = FONT_BODY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    insights = [
        ("Uneven Readiness", "In-person training attendance did not correlate with material distribution, causing operational drag."),
        ("Qwen Multilingual Synthesis", "Qwen organized Hindi & English observer notes in minutes, reducing manual coding time by 70%."),
        ("Actionable Policy Decision", "Suggested Response: Pair all future training sessions with mandatory digital/physical kits and 14-day follow-up checks."),
        ("Repeatable Approach", "This exact analytical model can be deployed across WhatsApp Chatbot and Assess Clear projects immediately.")
    ]

    for title, desc in insights:
        p_t = tf2.add_paragraph()
        p_t.text = f"\n🔹  {title}"
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        p_d = tf2.add_paragraph()
        p_d.text = f"     {desc}"
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: Empirical Proof 2: CRO Field Notes + Quant Data
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Empirical Proof #2", "CRO Study: Combining Hard Numbers with Field Notes", "Synthesizing 411 sampled classrooms across 55 districts using Qwen qualitative AI coding.")

    stats = [
        ("1,268", "Field Note Entries Grouped", "Qwen categorized 1,268 qualitative observation entries into structured intervention needs."),
        ("52.5%", "Student Attendance Rate", "6,328 present out of 12,053 enrolled students observed during field visits."),
        ("3.6%", "Fully Aligned Lesson Plans", "Only 15 out of 411 observed classrooms had written plans fully aligned with the Teacher Guide.")
    ]

    card_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    card_top = Inches(2.0)
    card_h = Inches(3.2)

    for i, (stat, title, desc) in enumerate(stats):
        x = start_x + i * (card_w + gap)
        add_card(s7, x, card_top, card_w, card_h)

        tb = s7.shapes.add_textbox(x + Inches(0.25), card_top + Inches(0.3), card_w - Inches(0.5), card_h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stat
        p.font.name = FONT_BODY
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = ACCENT_PURPLE if i == 0 else ACCENT_INDIGO

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    add_card(s7, Inches(0.8), Inches(5.4), Inches(11.733), Inches(1.5), bg_color=SURFACE_CARD, border_color=ACCENT_EMERALD)
    tb_b = s7.shapes.add_textbox(Inches(1.1), Inches(5.5), Inches(11.133), Inches(1.3))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "PRACTICAL ANALYTICAL TAKEAWAY"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD

    p_d = tf_b.add_paragraph()
    p_d.text = "By reading development needs (qualitative field notes coded by Qwen) alongside attendance and lesson alignment numbers, team members get the complete picture: Low lesson alignment (3.6%) was driven by lack of practical lesson preparation support, not teacher willingness."
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(12)
    p_d.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 8: Streamlit Analytics App
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Analytical Interface", "Streamlit Dashboard: Democratizing Data Access", "An intuitive visual portal allowing team members to explore findings interactively.")

    features = [
        ("01 / Choose", "Select any project, dataset, or research question from a unified dropdown menu without touching code."),
        ("02 / Explore", "Filter charts dynamically by district, teacher cohort, or observation period in real-time."),
        ("03 / Review Evidence", "Click any high-level stat to inspect the raw qualitative field notes and Qwen AI coding audit trail behind it."),
        ("04 / Download", "Export pre-approved summary reports, presentation-ready charts, and structured CSV tables instantly.")
    ]

    card_w = Inches(5.66)
    card_h = Inches(2.25)
    coords = [
        (Inches(0.8), Inches(2.0)),
        (Inches(6.86), Inches(2.0)),
        (Inches(0.8), Inches(4.55)),
        (Inches(6.86), Inches(4.55))
    ]

    for i, (title, desc) in enumerate(features):
        x, y = coords[i]
        add_card(s8, x, y, card_w, card_h)

        tb = s8.shapes.add_textbox(x + Inches(0.3), y + Inches(0.3), card_w - Inches(0.6), card_h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 9: Governance, Guardrails & Human Control
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Governance & Trust", "Guardrails: Maintaining Rigor, Ethics & Control", "AI accelerates research, but human expertise owns every final insight.")

    gov_cards = [
        ("AI Supports, Humans Own", "Qwen & DeepSeek act as first-pass copilots. Senior researchers verify meaning, cross-check context, and sign off on all team deliverables.", ACCENT_EMERALD),
        ("Full Auditability", "Every AI-generated theme links back to the exact source field note line. No black-box summaries—100% verifiable evidence trail.", ACCENT_INDIGO),
        ("Data Privacy & Security", "All data is anonymized prior to processing. Local open-weight models (Qwen, DeepSeek) prevent third-party data leakage.", ACCENT_PURPLE)
    ]

    card_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    card_top = Inches(2.0)
    card_h = Inches(4.8)

    for i, (title, desc, color) in enumerate(gov_cards):
        x = start_x + i * (card_w + gap)
        add_card(s9, x, card_top, card_w, card_h, border_color=BORDER_CARD)

        tb = s9.shapes.add_textbox(x + Inches(0.3), card_top + Inches(0.4), card_w - Inches(0.6), card_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"0{i+1}"
        p.font.name = FONT_BODY
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = color

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 10: 4-Phase Roadmap & Next Steps
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Execution Roadmap", "Strategic 4-Phase Rollout Plan", "Clear, low-risk milestones to validate accuracy before full deployment.")

    phases = [
        ("Phase 1: Prototype", "Immediate", "Build prototype pipeline using anonymized EDS data. Establish Qwen RAG context definitions & baseline speed metrics."),
        ("Phase 2: Historical Audit", "Month 2", "Re-analyze past projects (CLSS & CM RISE studies). Benchmark Qwen + TabFM against original human findings to prove accuracy."),
        ("Phase 3: Active Scaling", "Months 3-4", "Deploy pipeline across active research projects. Standardize research SOPs & launch Streamlit dashboard."),
        ("Phase 4: Full Deployment", "Months 5-6", "Evaluate long-term impact on research quality. Scale AI pipeline across all research & analytics departments.")
    ]

    card_w = Inches(2.75)
    gap = Inches(0.24)
    start_x = Inches(0.8)
    card_top = Inches(2.0)
    card_h = Inches(4.8)

    for i, (timeline, title, desc) in enumerate(phases):
        x = start_x + i * (card_w + gap)
        color = ACCENT_EMERALD if i == 0 else (ACCENT_PURPLE if i == 1 else ACCENT_INDIGO)
        add_card(s10, x, card_top, card_w, card_h, border_color=BORDER_CARD)

        tb = s10.shapes.add_textbox(x + Inches(0.2), card_top + Inches(0.3), card_w - Inches(0.4), card_h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title.upper()
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color

        p_t = tf.add_paragraph()
        p_t.text = timeline
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 11: Program Leadership & Team Credits
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Team Credits & Governance", "Program Leadership & Architectural Credits", "Recognizing the vision, leadership, and technical execution behind the AI stack.")

    credits_cards = [
        ("Concept & Program Context", "Shirish & Shil", "Shirish presented the original concept and program-level code design. Shil helped establish and shape the program context and analytical framework.", ACCENT_AMBER),
        ("Guidance & Leadership", "Rohan", "Provided continuous strategic guidance, leadership, and project oversight throughout the program lifecycle.", ACCENT_INDIGO),
        ("Lead Development & Execution", "Ashish", "Lead developer responsible for full-stack implementation, model integration, interactive dashboard development, and deployment.", ACCENT_EMERALD)
    ]

    card_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    card_top = Inches(2.0)
    card_h = Inches(4.8)

    for i, (role, name, desc, color) in enumerate(credits_cards):
        x = start_x + i * (card_w + gap)
        add_card(s11, x, card_top, card_w, card_h, border_color=BORDER_CARD)

        tb = s11.shapes.add_textbox(x + Inches(0.3), card_top + Inches(0.4), card_w - Inches(0.6), card_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = role.upper()
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color

        p_t = tf.add_paragraph()
        p_t.text = name
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 12: Integrated Live Analytics Dashboards (4 Portals)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Live Dashboards & Portals", "Integrated Ecosystem of 4 Live Analytics Portals", "Direct cloud & web access to operational dashboards and research telemetry.")

    dash_cards = [
        ("EDS Streamlit Dashboard", "https://cm-rise-eds-dashboard.streamlit.app/", "Teacher training material availability & field observation telemetry.", ACCENT_ROSE),
        ("CRO Streamlit Dashboard", "https://peepul-mp-cpd-cro-dashboard.streamlit.app/", "Classroom observation & attendance alignment analytics across 411 schools.", ACCENT_PURPLE),
        ("Lifted Analytics Dashboard", "https://ashishghanghoriya1-png.github.io/lifted-analytics-dashboard/", "Intervention fidelity & multi-district telemetry infrastructure.", ACCENT_INDIGO),
        ("MP Shaikshik Samwaad", "https://ashishghanghoriya1-png.github.io/MP-Shaikshik-Samwaad/", "Academic dialogue metrics & teacher engagement analytics across Madhya Pradesh.", ACCENT_EMERALD)
    ]

    card_w = Inches(2.75)
    gap = Inches(0.24)
    start_x = Inches(0.8)
    card_top = Inches(2.0)
    card_h = Inches(4.8)

    for i, (title, url, desc, color) in enumerate(dash_cards):
        x = start_x + i * (card_w + gap)
        add_card(s12, x, card_top, card_w, card_h, border_color=BORDER_CARD)

        tb = s12.shapes.add_textbox(x + Inches(0.2), card_top + Inches(0.3), card_w - Inches(0.4), card_h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"PORTAL 0{i+1}"
        p.font.name = FONT_BODY
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = color

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc + "\n\nURL: " + url
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MUTED

    output_path = "AI_Research_Analytics_Executive_Briefing.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully in Strategic Research Briefing style: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_deck()
